from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PInches, Pt as PPt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from .physics_diagrams import render_physics_diagram_svg


def _set_doc_font(doc: Document) -> None:
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(14 if name == "Normal" else 16)


def create_docx(title: str, content: str) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(0.8)
    _set_doc_font(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold, run.font.name, run.font.size = True, "Times New Roman", Pt(16)
    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith(("- ", "• ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    out = BytesIO(); doc.save(out); return out.getvalue()


def create_pdf(title: str, content: str) -> bytes:
    out = BytesIO()
    font = "Helvetica"
    candidates = [Path("C:/Windows/Fonts/times.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf")]
    for path in candidates:
        if path.exists():
            pdfmetrics.registerFont(TTFont("KZSerif", str(path))); font = "KZSerif"; break
    styles = getSampleStyleSheet()
    for s in styles.byName.values():
        s.fontName = font
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for line in content.splitlines():
        if line.strip(): story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), styles["BodyText"])); story.append(Spacer(1, 5))
    SimpleDocTemplate(out, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42).build(story)
    return out.getvalue()


def create_pptx(title: str, slides: list[dict[str, Any]], exact_count: int | None = None) -> bytes:
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    wanted = exact_count or len(slides) or 7
    normalized = slides[:wanted]
    while len(normalized) < wanted:
        normalized.append({"title": "Қорытынды" if len(normalized) == wanted - 1 else f"{len(normalized)+1}-слайд", "bullets": []})
    for i, item in enumerate(normalized):
        slide = prs.slides.add_slide(prs.slide_layouts[0] if i == 0 else prs.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title") or (title if i == 0 else f"{i+1}-слайд"))
        body = slide.placeholders[1]
        bullets = item.get("bullets") or ([str(item.get("subtitle") or "")] if i == 0 else [])
        body.text = "\n".join(str(x) for x in bullets)
        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            for p in shape.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"; r.font.size = PPt(24 if shape == slide.shapes.title else 19)
    out = BytesIO(); prs.save(out); return out.getvalue()


def create_diagram_svg(title: str, labels: list[str]) -> bytes:
    nodes = (labels or ["Берілгені", "Формула", "Есептеу", "Жауап"])[:6]
    width, height = 1200, max(420, 145 * len(nodes))
    items = []
    for i, label in enumerate(nodes):
        y = 60 + i * 125
        items.append(f'<rect x="210" y="{y}" width="780" height="78" rx="22" fill="#eef5ff" stroke="#1769ff" stroke-width="3"/>')
        safe = str(label).replace("&", "&amp;").replace("<", "&lt;")
        items.append(f'<text x="600" y="{y+49}" text-anchor="middle" font-family="Times New Roman" font-size="28" fill="#10213d">{safe}</text>')
        if i < len(nodes)-1:
            items.append(f'<path d="M600 {y+78} L600 {y+119}" stroke="#1769ff" stroke-width="4" marker-end="url(#a)"/>')
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" role="img"><defs><marker id="a" markerWidth="12" markerHeight="12" refX="10" refY="6" orient="auto"><path d="M0,0 L12,6 L0,12 z" fill="#1769ff"/></marker></defs><rect width="100%" height="100%" fill="white"/><text x="600" y="38" text-anchor="middle" font-family="Times New Roman" font-size="30" font-weight="bold">{title}</text>{''.join(items)}</svg>'''
    return svg.encode("utf-8")


def create_physics_diagram_svg(title: str, spec: dict[str, Any] | None = None) -> bytes:
    """Render a deterministic textbook-style physics diagram from structured data."""
    return render_physics_diagram_svg(spec or {}, title=title)
