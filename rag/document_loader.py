from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
import pandas as pd


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    if suffix == ".pptx":
        prs = Presentation(BytesIO(data))
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text)
    if suffix == ".xlsx":
        sheets = pd.read_excel(BytesIO(data), sheet_name=None)
        return "\n\n".join(f"[{name}]\n{frame.to_csv(index=False)}" for name, frame in sheets.items())
    if suffix in {".txt", ".md", ".csv"}:
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Бұл файлдың мәтін қабатын оқу мүмкін болмады")


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 180) -> list[str]:
    clean = " ".join(text.split())
    if not clean:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + chunk_size)
        chunk = clean[start:end]
        if end < len(clean):
            last = max(chunk.rfind(". "), chunk.rfind("! "), chunk.rfind("? "))
            if last > chunk_size // 2:
                end = start + last + 1
                chunk = clean[start:end]
        chunks.append(chunk.strip())
        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)
    return [c for c in chunks if len(c) > 40]
