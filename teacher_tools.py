from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt


def _clean_markdown(text: str) -> str:
    text = re.sub(r"```.*?```", lambda m: m.group(0).replace("```", ""), text, flags=re.S)
    text = text.replace("**", "")
    return text.strip()


def build_docx_bytes(title: str, content: str) -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(14)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            styles[style_name].font.name = "Times New Roman"
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title.strip() or "Құжат")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for raw in _clean_markdown(content).splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^[-•]\s+", line):
            p = doc.add_paragraph(re.sub(r"^[-•]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            p = doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
        else:
            p = doc.add_paragraph(line)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def build_pptx_bytes(title: str, slides: list[dict[str, Any]]) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def set_shape_font(shape, size=24, bold=False):
        if not getattr(shape, "has_text_frame", False):
            return
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = PPTPt(size)
                run.font.bold = bold

    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title.strip() or "Презентация"
    first.placeholders[1].text = "AI Physics KZ · Мұғалім ЖИ ассистенті"
    set_shape_font(first.shapes.title, 30, True)
    set_shape_font(first.placeholders[1], 18, False)

    for item in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        heading = str(item.get("title") or "Слайд").strip()
        bullets = item.get("bullets") or item.get("content") or []
        if isinstance(bullets, str):
            bullets = [x.strip() for x in bullets.split("\n") if x.strip()]
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets[:8]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(bullet)
            p.level = 0
        set_shape_font(slide.shapes.title, 26, True)
        set_shape_font(slide.placeholders[1], 20, False)

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def infer_artifact_kind(prompt: str) -> str:
    low = prompt.lower()
    # «Презентацияға фон/сурет жаса» деген сұрау PPTX емес, сурет генерациясы болуы керек.
    if any(k in low for k in ["фон", "сурет", "иллюстрац", "постер", "image", "көрнекі"]):
        return "image"
    if any(k in low for k in ["презентац", "слайд", "ppt", "powerpoint"]):
        return "pptx"
    if any(k in low for k in ["қмж", "бжб", "тжб", "жұмыс пара", "word", "docx", "жоспар", "анықтама", "есеп беру"]):
        return "docx"
    return "text"
