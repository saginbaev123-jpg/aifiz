from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reports.learning_evidence import paired_diagnostics


def build_student_docx(user: dict[str, Any], mastery: list[dict[str, Any]], mistakes: list[dict[str, Any]], diagnostics: list[dict[str, Any]], learning_path: list[dict[str, str]]) -> bytes:
    doc = Document()
    styles = doc.styles
    for style in styles:
        if style.type == 1:  # paragraph styles, including headings
            style.font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(14)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI Physics KZ — оқушы есебі")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    doc.add_paragraph(f"Оқушы: {user.get('full_name','')}")
    doc.add_paragraph(f"Сынып: {user.get('grade','')}")

    doc.add_heading("Тақырыптық меңгеру", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Тақырып"
    hdr[1].text = "Меңгеру"
    hdr[2].text = "Әрекет"
    hdr[3].text = "Дұрыс"
    for r in mastery:
        cells = table.add_row().cells
        cells[0].text = str(r["topic"])
        cells[1].text = f"{float(r.get('score',0)):.0f}%"
        cells[2].text = str(r.get("attempts",0))
        cells[3].text = str(r.get("correct",0))

    doc.add_heading("Диагностика динамикасы", level=1)
    if diagnostics:
        for d in diagnostics[:5]:
            doc.add_paragraph(f"{d['created_at'][:10]} — {d['correct']}/{d['total']} ({d['percent']:.1f}%)")
    else:
        doc.add_paragraph("Диагностика әлі орындалмаған.")
    doc.add_heading("Салыстырмалы оқу нәтижесі", level=1)
    pairs = paired_diagnostics(diagnostics)
    if pairs:
        for item in pairs:
            doc.add_paragraph(
                f"{item['Тақырып']}: {item['Бастапқы']:.1f}% → {item['Қорытынды']:.1f}%, "
                f"өзгеріс {item['Өзгеріс (п.т.)']:+.1f} пайыздық тармақ "
                f"({item['Сұрақ саны']} сұрақ, {item['Бастапқы күні']} – {item['Қорытынды күні']})."
            )
    else:
        doc.add_paragraph("Салыстыруға жарамды бастапқы және қорытынды диагностика әлі жоқ.")

    doc.add_heading("Қайталанатын қателер", level=1)
    if mistakes:
        for m in mistakes[:10]:
            doc.add_paragraph(f"{m['topic']}: {m['mistake_type']} — {m['description']} ({m['frequency']} рет)")
    else:
        doc.add_paragraph("Қайталанатын қате тіркелмеген.")

    doc.add_heading("Жеке оқу траекториясы", level=1)
    for p in learning_path:
        doc.add_paragraph(f"{p['topic']} ({p['score']}): {p['action']}")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
